import os


def extract_resume_text(file_path, file_format):
    """Extract raw text from a resume file based on its format."""
    file_format = (file_format or '').lower().lstrip('.')

    if file_format == 'pdf':
        return _extract_pdf(file_path)
    elif file_format in ('docx', 'doc'):
        return _extract_docx(file_path)
    elif file_format == 'txt':
        return _extract_txt(file_path)
    else:
        # Try fallback detection based on extension
        ext = os.path.splitext(file_path)[1].lower().lstrip('.')
        if ext == 'pdf':
            return _extract_pdf(file_path)
        elif ext in ('docx', 'doc'):
            return _extract_docx(file_path)
        else:
            return _extract_txt(file_path)


def _extract_pdf(file_path):
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return '\n\n'.join(text_parts)
    except Exception as e:
        raise Exception(f"Failed to extract text from PDF: {str(e)}")


def _extract_docx(file_path):
    try:
        from docx import Document
        doc = Document(file_path)
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(' | '.join(cells))
        return '\n'.join(parts)
    except Exception as e:
        raise Exception(f"Failed to extract text from DOCX: {str(e)}")


def _extract_txt(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    except Exception as e:
        raise Exception(f"Failed to read text file: {str(e)}")
